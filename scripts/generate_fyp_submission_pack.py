from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\kg_sma_0704")
FYP = Path(r"C:\Users\jon15\Desktop\大四上\FYP")
TEMP = ROOT / "tmp" / "fyp_documents_2026-10-01" / "authoring"
OUT = ROOT / "outputs" / "fyp_submission_pack_2026-10-01"

PROJECT_TITLE = "Evidence-Grounded Construction of a Knowledge Graph for Spinal Muscular Atrophy"
STUDENT = "贾欧妮 [insert official English name]"
IDS = "UoG: [TO CONFIRM]    UESTC: [TO CONFIRM]"
SUPERVISOR = "[TO CONFIRM]"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def style_run(run, size=10.5, bold=False, color=None, name="Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_doc_defaults(doc: Document, font="Arial", size=10.5, line=1.08) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.font.size = Pt(size)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = line
    for name, points, color in (
        ("Title", 20, (31, 78, 121)),
        ("Heading 1", 15, (31, 78, 121)),
        ("Heading 2", 12, (47, 84, 117)),
        ("Heading 3", 10.5, (47, 84, 117)),
    ):
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(points)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def set_style_shading(style, fill: str | None) -> None:
    p_pr = style.element.get_or_add_pPr()
    existing = p_pr.find(qn("w:shd"))
    if existing is not None:
        p_pr.remove(existing)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        p_pr.append(shd)


def style_preliminary_headings(doc: Document) -> None:
    h1 = doc.styles["Heading 1"]
    h1.font.color.rgb = RGBColor(255, 255, 255)
    set_style_shading(h1, "2F5E85")
    h1.paragraph_format.left_indent = Cm(0.12)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(4)

    h2 = doc.styles["Heading 2"]
    h2.font.color.rgb = RGBColor(31, 78, 121)
    set_style_shading(h2, "D9EAF7")
    h2.paragraph_format.left_indent = Cm(0.10)
    h2.paragraph_format.space_before = Pt(5)
    h2.paragraph_format.space_after = Pt(3)


def configure_a4(doc: Document, margins=(1.8, 1.8, 1.7, 1.7)) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(margins[0])
        section.bottom_margin = Cm(margins[1])
        section.left_margin = Cm(margins[2])
        section.right_margin = Cm(margins[3])


def add_footer_page_number(doc: Document) -> None:
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Page ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)


def add_banner(doc: Document, text: str, fill="D9EAF7", color=(31, 78, 121)) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=90, bottom=90)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    style_run(run, 10, True, color)


def add_bullets(doc: Document, items: list[str], level=0, size=10.5) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55 + 0.4 * level)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(("• " if level == 0 else "– ") + item)
        style_run(run, size)


def add_numbered(doc: Document, items: list[str], size=10.5) -> None:
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"{idx}. {item}")
        style_run(run, size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, "2F5E85")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        style_run(r, font_size, True, (255, 255, 255))
        if widths:
            cell.width = Cm(widths[i])
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cells[i], "F3F6F8")
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            r = p.add_run(value)
            style_run(r, font_size)
            if widths:
                cells[i].width = Cm(widths[i])
    return table


def add_para(doc: Document, text: str, bold_lead: str | None = None, size=10.5, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        style_run(r1, size, True)
        r2 = p.add_run(text[len(bold_lead):])
        style_run(r2, size)
    else:
        style_run(p.add_run(text), size)
    return p


def make_roadmap() -> Path:
    doc = Document()
    configure_a4(doc)
    set_doc_defaults(doc, font="Microsoft YaHei", size=10.5, line=1.12)
    add_footer_page_number(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("FYP 进度、提交与写作指南"), 20, True, (31, 78, 121), "Microsoft YaHei")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(PROJECT_TITLE), 12, True, (47, 84, 117), "Arial")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("编制日期：2026-10-01｜计划起点：2026-10-01｜学年：2026/27"), 9.5, False, (80, 80, 80), "Microsoft YaHei")
    add_banner(doc, "真实性原则：本计划从现在开始安排未来工作；不把现有成果伪造成过去逐周完成的 logbook 记录。")

    doc.add_heading("1. 已确认的提交节点", level=1)
    add_table(
        doc,
        ["提交/考核", "权重", "截止时间（CST）", "交付内容与渠道"],
        [
            ["Specification & Preliminary Report", "10%", "2026-10-30 23:59", "Preliminary Report + 第1本 Logbook + Risk Assessment；GC-UESTC FYP System"],
            ["Interim Report + Oral + 第2本 Logbook", "15%", "报告：2027-01-08 23:59\n口试：2027-01-11—22", "FYP System；报告、演示与问答综合考核"],
            ["Poster submission", "—", "2027-04-16 12:00", "FYP System"],
            ["Poster oral presentation", "25%", "2027-04-19—21", "海报讲解与现场问答"],
            ["Final Report + 第3本 Logbook", "40%", "时间存在冲突，见下方警告", "纸质提交 + Moodle + FYP System"],
            ["Student Performance Evaluation", "10%", "2027-04-28", "导师通过 FYP System 评价"],
            ["Software/code/prototype", "—", "时间存在冲突，见下方警告", "软件代码与相关材料归档提交"],
        ],
        widths=[4.4, 1.5, 4.2, 7.0],
        font_size=8.7,
    )
    add_banner(
        doc,
        "截止时间冲突：FYP_timeline.png 写明 Final Report、第三本 Logbook 和代码为 2027-04-28 23:59；Student Handbook 2026–27 表 1 写明 Final Report/Logbook 与代码为 2027-04-30 23:59。请在 2027 年 4 月前以 Moodle、FYP System 和最新 assessment brief 为准，并按较早的 4 月 28 日做内部截止。",
        fill="FCE4D6",
        color=(156, 63, 28),
    )

    doc.add_heading("2. 从零开始的未来执行计划", level=1)
    add_para(doc, "安排原则：每阶段都有可检查的输入、任务、输出和验证证据；只有真正完成后才在 logbook 中写成完成项。现有代码与数据可作为技术蓝图和验证基线，但日志必须记录你从 2026-10-01 起亲自执行、理解、修改和评估的工作。", size=10)
    roadmap_rows = [
        ["10/01–10/04", "立项与合规", "阅读指南；确认研究问题、范围、风险、伦理、资源和评价指标", "签字版 risk/ethical；项目问题清单；第1周真实日志"],
        ["10/05–10/11", "文献与需求", "SMA、PubMed/Open Targets、BioNLP、KG、RAG 文献检索；定义实体/关系 schema", "文献矩阵；schema v1；可测量目标"],
        ["10/12–10/18", "数据采集", "配置环境；实现/复现 PubMed 与 Open Targets 采集；数据质量检查", "原始数据、manifest、计数、哈希、失败记录"],
        ["10/19–10/25", "主题分析与抽取小样", "BERTopic/PubMedBERT 主题分析；LLM 关系抽取提示词与小规模 pilot", "主题报告；pilot triples；错误分类"],
        ["10/26–10/30", "首轮提交", "完善 preliminary report、风险、伦理、第1本 logbook；导师签字与提交检查", "10/30 提交包；后续修订清单"],
        ["11/01–11/15", "全量关系抽取", "分块、并发、重试、schema 校验、去重；保存 run artifacts", "可复现的全量 triples 与 validation summary"],
        ["11/16–11/30", "标准化与语义对齐", "词典映射、同义词归一、PubMedBERT 类型内对齐；阈值敏感性分析", "mapped/aligned triples；阈值实验"],
        ["12/01–12/15", "融合与冲突处理", "聚合证据、置信度融合、正负关系冲突检测、人工复核规则", "fused triples；conflict list；错误案例"],
        ["12/16–12/31", "图谱与中期材料", "Neo4j 导入、拓扑指标、可视化；撰写 interim report 和演示", "图数据库、viewer、指标；中期报告草稿"],
        ["01/01–01/08", "中期提交", "复现核心结果、修订 Gantt、完成报告与第2本 logbook", "1/8 提交包"],
        ["01/11–01/22", "中期口试", "准备 5–10 分钟技术叙述、风险/局限、问答；记录反馈", "演示稿；问答库；反馈行动项"],
        ["01/23–02/14", "400条人工标注", "300 条随机主集 + 100 条困难集；80 条独立双评；计算一致性", "人工 gold set；Cohen’s κ；precision 与错误分析"],
        ["02/15–03/07", "创新1：证据校验", "把 evidence span 对齐回 PubMed title/abstract；设置阈值与人工复核路由", "代码、测试、覆盖率/准确率、失败案例"],
        ["03/08–03/21", "创新2：GraphRAG", "Neo4j 邻域检索、PMID 证据组装、回答生成与拒答策略", "端到端 runner；QA 测试集；基线对比"],
        ["03/22–04/04", "创新3–4：引用与主张校验", "验证 PMID 引用存在且被检索；拆分 claims 并判断证据支持度", "citation/claim metrics；消融；错误分析"],
        ["04/05–04/15", "收尾与海报", "冻结实验、重跑主结果、制作海报、导师审阅", "最终表图；4/16 海报"],
        ["04/19–04/21", "海报口试", "展示贡献、证据链、局限与未来工作；记录面板反馈", "口试记录；最终报告修订项"],
        ["04/22–04/28", "最终内部截止", "完成论文、第三本 logbook、代码与复现包；查重和引用核对", "按 4/28 内部截止完成全部材料"],
        ["04/29–04/30", "仅作缓冲", "仅在官方系统确认 4/30 截止时用于上传/技术问题，不安排新实验", "最终系统回执"],
    ]
    add_table(doc, ["日期", "工作包", "主要任务", "完成证据"], roadmap_rows, widths=[2.5, 3.2, 7.2, 4.9], font_size=8.1)

    doc.add_page_break()
    doc.add_heading("3. 每一阶段怎样写进 logbook", level=1)
    add_para(doc, "Logbook 是当时发生的工程活动记录，不是事后整理的漂亮报告。建议每周 5–6 页有思考过程的记录；每项工作写日期、开始/结束时间、目标、方法、命令或草图、观察、问题、决策依据、结果、下一步，并给每页编号和索引。错误和失败要保留，错误内容用单线划掉后写更正，不能撕页、涂改或回填。", size=10.5)
    add_table(
        doc,
        ["活动", "应记录内容", "本项目示例"],
        [
            ["文献检索", "检索式、数据库、筛选标准、保留/排除理由、获得的结论", "为何选择 PubMed/Open Targets；哪些工作定义了 relation schema"],
            ["编码/实验", "环境、版本、输入、命令、参数、开始/结束时间、原始输出、异常", "LLM 批次、worker 数、失败重试、数据哈希、Neo4j 导入计数"],
            ["分析/决策", "候选方案、评价标准、证据、最终选择和替代方案", "为何只在同一实体类型内做 embedding alignment；阈值如何选"],
            ["会议", "会前议题、讨论结论、待办、负责人、期限、下次会议", "不要填不存在的会议；邮件和线上会议数量按实际记录"],
            ["结果", "表图、单位、统计量、解释、局限、与预期不一致之处", "精度只用 300 条随机集；100 条困难集单独报告；不能由 precision 推断 recall"],
            ["创新", "动机、基线、设计、实现、测试、量化比较、失败案例", "证据校验、GraphRAG、引用校验、claim–evidence 校验"],
        ],
        widths=[2.4, 7.0, 8.4],
        font_size=8.8,
    )
    doc.add_heading("首本 logbook 的当前状态", level=2)
    add_bullets(doc, [
        "本交付包中的第一本 logbook 只记录 2026-10-01 已实际完成的指南阅读、材料审查、范围确认和表格草拟。",
        "10 月后续日期提供的是“计划提示”，必须在每天实际工作后用真实记录替换；不能直接作为已完成记录提交。",
        "封面中的会议次数、会议方式和邮件数量保持 [TO CONFIRM]，因为材料没有提供事实依据。",
        "每月封面是摘要；提交时还需扫描/合并对应时间段的手写或原始电子 logbook 页面。",
    ], size=10)

    doc.add_heading("4. 报告体系与写法", level=1)
    add_table(
        doc,
        ["报告", "核心目的", "建议结构与证据"],
        [
            ["Preliminary Report", "证明问题值得做、目标可测、路线可行", "Project Description；Measurable Outcomes；≥4页 Technical Background；References；≤1页 Work Plan；总计≤8页"],
            ["Interim Report", "证明已有进展、能批判性解释结果，并给出修订计划", "Project Outline；Work Done So Far；初步结论；Challenges/Solutions；剩余工作；Revised Gantt；References；≤12页"],
            ["Final Dissertation", "完整呈现背景、可复现方法、结果、解释、贡献和局限", "Title/Declaration；Abstract；Acknowledgements；Contents；Introduction；Literature Review；Methods；Results；Discussion；Conclusions/Future Work；References；必要附录"],
            ["Poster + oral", "用最少信息说明问题、方法、最重要的 1–2 个结果和原创贡献", "一张主流程图、2–4 个关键结果图、明确创新、局限和可追溯证据；准备经典问题"],
        ],
        widths=[3.4, 5.4, 9.0],
        font_size=8.7,
    )
    add_bullets(doc, [
        "Introduction 解释为什么做、研究问题是什么；Literature Review 从广到窄并批判已有工作，最后落到 research gap。",
        "Methods 要让他人能够复现：数据来源、检索日期、schema、模型/版本、提示词、参数、质量门、统计方法都要写清。",
        "Results 只陈述观察结果，表图编号一致、坐标/单位/误差清楚；Discussion 解释原因、对照文献、说明局限。",
        "Abstract 最后写，概括问题、方法、最主要结果和结论；不设置小标题也不塞入新信息。",
        "图题置于图下，表题置于表上；复制或重绘他人材料都要标注来源。引用格式必须前后一致。",
        "严禁购买报告、拼接他人文字、近似改写却不引用，或把生成式 AI 输出当作事实。使用 AI 时遵守 assessment brief，核对事实与参考文献并按要求声明。",
    ], size=10)

    doc.add_heading("5. Risk、Ethics、元器件与报销结论", level=1)
    add_para(doc, "Risk Assessment：", bold_lead="Risk Assessment：", size=10.5)
    add_para(doc, "需要提交并由学生与第一导师签字。主要风险不是化学/机械危害，而是长期屏幕工作、数据和代码丢失、密钥泄露、第三方 API 成本与限流、LLM 幻觉、引用/知识产权和错误医学解释。对应控制包括休息与人体工学、版本控制与备份、.env 密钥管理、预算/重试/检查点、schema 与证据校验、人工标注和明确“非临床建议”。")
    add_para(doc, "Ethical Declaration：", bold_lead="Ethical Declaration：", size=10.5)
    add_para(doc, "需要提交。当前设计仅处理公开 PubMed 摘要与 Open Targets 公共数据，不招募受试者、不访问病历、不处理可识别患者信息、不进行生物样本或动物实验，因此 12 个 gateway questions 均为 No。若以后加入患者数据、访谈或私有临床文本，必须重新申请伦理审查。")
    add_para(doc, "Component / Reimbursement：", bold_lead="Component / Reimbursement：", size=10.5)
    add_para(doc, "当前不需要填写。两份样表用于元器件采购或发票报销，而本项目可用现有计算机、开源软件、公开数据和既有 API 配额完成。只有在导师事先批准购买硬件、服务器/GPU、数据库服务或可报销 API 费用时，才应填写，并使用真实供应商、单价、发票和银行信息；不得为了“完整”而虚构采购。")

    doc.add_heading("6. 400 条人工标注是否必要", level=1)
    add_para(doc, "结论：作为毕业设计，固定“400 条”并非学校硬性规定，但独立人工标注对这个项目是必要的。没有 gold set，只能报告条数和图拓扑，不能有说服力地报告关系抽取的事实正确性。400 条规模在时间与可信度之间较合理。")
    add_table(
        doc,
        ["集合", "数量", "抽样与用途", "报告规则"],
        [
            ["Primary random", "300", "从预测三元组中可复现随机抽样", "只用这 300 条计算主 precision/支持率与置信区间"],
            ["Challenge", "100", "过采样低置信、非精确证据、稀有关系", "仅作压力测试与错误分析，不能与随机集混算主指标"],
            ["Second review", "80", "从两组中预先标记，第二评审独立复核", "报告 Cohen’s κ/一致率；分歧经讨论裁决"],
        ],
        widths=[3.0, 1.8, 7.0, 6.0],
        font_size=9.0,
    )
    add_bullets(doc, [
        "每条先读 title + abstract，再核对 subject、subject type、relation、object、object type、方向和 evidence span。",
        "Support label：2=直接支持；1=部分/隐含支持；0=不支持或矛盾；U=摘要不足以判断。另填错误类型与修正三元组。",
        "AI 可给建议，但最终 gold label 必须由人确认；不能把 AI 建议伪装成人工标注。",
        "由于抽样起点是预测三元组，这套标注估计 precision/事实支持度，不估计 recall。Recall 需要对一批完整摘要穷尽标注所有关系。",
    ], size=10)

    doc.add_heading("7. 提交前检查清单", level=1)
    add_numbered(doc, [
        "用官方英文姓名、UoG/UESTC 学号、课程代码、Degree Programme 和 Supervisor 替换全部 [TO CONFIRM]。",
        "学生与导师完成 Risk Assessment 签字；Ethical Declaration 按学院要求签字。",
        "把 logbook 封面会议次数换成真实数字，并附对应日期的原始页面；删除所有尚未发生的 planned prompts。",
        "核对 Preliminary Report 总页数≤8、Technical Background≥4页、References 完整且所有引用可追溯。",
        "用 Turnitin/学校系统检查相似度；逐条核验生成式 AI 辅助文字和参考文献；保存 AI 使用声明（如 assessment brief 要求）。",
        "在 Moodle/FYP System 再确认最终截止日期；上传后保存回执和文件哈希。",
    ], size=10)

    doc.add_heading("材料依据", level=2)
    add_bullets(doc, [
        "A Guide to FYP Logbook.pdf",
        "Student Handbook FYP 2026-27.pdf",
        "Project writing and plagiarism printing 2017.pdf",
        "FYP_timeline.png",
        "毕设总体结构-贾欧妮.pdf 与 D:\\kg_sma_0704 的代码、run artifacts 和复现文档",
    ], size=9.5)

    path = OUT / "01_FYP_Roadmap_and_Submission_Guide.docx"
    doc.save(path)
    return path


def _replace_cell(cell, text: str, size=9.0, bold=False, color=None, compact=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if compact:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    style_run(p.add_run(text), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if compact:
        set_cell_margins(cell, top=0, start=20, bottom=0, end=20)
    else:
        set_cell_margins(cell)


def collapse_bordered_paragraph(paragraph) -> None:
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    run = paragraph.add_run(" ")
    style_run(run, 1)


def make_risk() -> Path:
    src = FYP / "Project Risk Assessment Form (CDHN10.30).docx"
    doc = Document(src)
    t0 = doc.tables[0]
    _replace_cell(t0.cell(0, 1), STUDENT, 9.0)
    _replace_cell(t0.cell(1, 1), IDS, 8.5)
    _replace_cell(t0.cell(2, 1), PROJECT_TITLE, 8.5)
    _replace_cell(t0.cell(3, 1), SUPERVISOR, 9.0)

    work = (
        "Software research project to construct an evidence-grounded knowledge graph for spinal muscular atrophy (SMA). "
        "The work will retrieve public PubMed abstracts and Open Targets records; cluster literature topics; extract and validate biomedical triples with NLP/LLM methods; normalise and fuse entities/relations; import the graph into Neo4j; and evaluate extraction quality, graph topology and evidence-grounded retrieval. "
        "No laboratory work, human recruitment, private clinical records, biological samples or animal work is planned."
    )
    doc.paragraphs[10].text = work
    style_run(doc.paragraphs[10].runs[0], 8.5)
    for idx in range(11, 19):
        collapse_bordered_paragraph(doc.paragraphs[idx])

    hazards = (
        "1. Prolonged computer use: eye strain, fatigue and musculoskeletal discomfort. Controls: ergonomic workstation, regular breaks and sensible working hours.\n"
        "2. Data/code loss or non-reproducible results. Controls: Git version control, dated run artefacts, manifests/hashes and backed-up outputs.\n"
        "3. Credential/cybersecurity risk. Controls: store API keys and Neo4j credentials only in ignored .env files; never place secrets in code, reports or logs.\n"
        "4. LLM hallucination or incorrect biomedical interpretation. Controls: restricted schema, evidence linking, validation gates, human review of a 400-item sample, and clear statement that outputs are not clinical advice.\n"
        "5. Third-party API cost, rate limits and service failure. Controls: approved budgets, quotas, retries, checkpointing and pilot runs.\n"
        "6. Copyright/licensing and academic integrity. Controls: use permitted public metadata/abstracts, cite sources, avoid redistributing full copyrighted text, and disclose permitted GenAI assistance.\n"
        "7. Neo4j/data overwrite risk. Controls: export/backup before imports; use scoped scripts and verify target database and counts before destructive operations."
    )
    doc.paragraphs[21].text = hazards
    style_run(doc.paragraphs[21].runs[0], 8.0)
    for idx in range(22, 29):
        collapse_bordered_paragraph(doc.paragraphs[idx])

    t1 = doc.tables[1]
    _replace_cell(
        t1.cell(0, 0),
        "Will the student receive relevant safety training?   ☒ Yes   ☐ No\n"
        "Details: project induction; Python/Conda and Git; data protection; safe API/secret handling; Neo4j backup/import; academic integrity and permitted GenAI use; ergonomic workstation practice.\n"
        "Provider: first supervisor, University guidance and relevant software documentation.\n"
        "Is training adequate to control identified hazards?   ☒ Yes   ☐ No",
        8.2,
    )
    _replace_cell(doc.tables[2].cell(0, 0), "Will the COSHH regulations apply?   ☐ Yes   ☒ No\nNo chemicals, biological agents or laboratory substances will be used.", 8.5)
    _replace_cell(doc.tables[2].cell(0, 1), "Not applicable.", 8.5)
    _replace_cell(doc.tables[3].cell(0, 0), "Will the student be supervised at all times?\nNo. The work is desk-based and low risk. Regular progress meetings and escalation to the first supervisor will be used.", 8.5)
    _replace_cell(doc.tables[3].cell(0, 1), "☐ Yes   ☒ No\nSupervisor: " + SUPERVISOR, 8.5)
    _replace_cell(
        doc.tables[4].cell(0, 0),
        "Are you satisfied that risks are adequately controlled?   ☒ Yes   ☐ No\n"
        "Controls will be reviewed if the project scope changes, particularly if private clinical data, human participants, new paid services or laboratory work are introduced.",
        8.5,
    )
    for row in doc.tables[5].rows:
        _replace_cell(row.cells[1], "☐ Yes   ☒ No", 7.5, compact=True)
    _replace_cell(
        doc.tables[6].cell(0, 0),
        "The project analyses public PubMed abstracts and Open Targets records only. It does not recruit participants or use patient-level, identifiable or confidential data. "
        "A second reviewer may label public-text extraction examples; only a minimal reviewer code will be recorded. LLM outputs may be inaccurate, so claims will be linked to PMID evidence, evaluated by human review and presented as research outputs rather than clinical advice. "
        "Copyright, source licensing, attribution, data provenance and permitted GenAI use will be documented. Any later change involving human participants, private clinical text or biological material requires a new ethical and risk review before work begins.",
        8.2,
    )
    sig = doc.tables[7]
    _replace_cell(sig.cell(0, 0), "Signed (student): [SIGNATURE REQUIRED]", 8.5)
    _replace_cell(sig.cell(0, 1), "Date: [TO CONFIRM]", 8.5)
    _replace_cell(sig.cell(1, 0), "Signed (1st Supervisor): [SIGNATURE REQUIRED]", 8.5)
    _replace_cell(sig.cell(1, 1), "Date: [TO CONFIRM]", 8.5)

    path = OUT / "02_Project_Risk_Assessment_Draft.docx"
    doc.save(path)
    return path


def make_ethics() -> Path:
    src = FYP / "Ethical Consideration Form10.30.docx"
    doc = Document(src)
    t0 = doc.tables[0]
    _replace_cell(t0.cell(0, 1), PROJECT_TITLE, 8.5)
    _replace_cell(t0.cell(1, 1), STUDENT, 9.0)
    _replace_cell(t0.cell(2, 1), IDS, 8.5)
    _replace_cell(t0.cell(3, 1), SUPERVISOR, 9.0)
    for row in doc.tables[1].rows:
        _replace_cell(row.cells[1], "☐ Yes   ☒ No", 7.5, compact=True)
    _replace_cell(
        doc.tables[2].cell(0, 0),
        "This is a desk-based software and secondary-data project. It will use public PubMed titles/abstracts and public Open Targets records. It does not recruit human participants; access patient records; collect names, contact details, audio/video or other identifiable information; use biological samples; conduct interventions; or involve animals.\n\n"
        "The principal ethical issues are accuracy, provenance, copyright, bias and the risk that generated biomedical statements could be mistaken for clinical advice. Controls are: preserve PMID/source provenance; cite all external material; use only permitted data; do not redistribute full copyrighted articles; store API keys in ignored local environment files; evaluate extraction correctness with a human-reviewed sample; label uncertainty and conflicts; and state that the graph/GraphRAG output is for research and not diagnosis or treatment.\n\n"
        "If a second reviewer helps annotate public abstracts, participation is a project contribution rather than a research subject procedure. Only a minimal reviewer code will be stored, with no unnecessary personal data. Any future use of private clinical data, interviews, patient groups or biological material will require fresh ethical review before collection begins.",
        7.7,
    )
    sig = doc.tables[3]
    _replace_cell(sig.cell(0, 0), "Signed (student): [SIGNATURE REQUIRED]", 8.5)
    _replace_cell(sig.cell(0, 1), "Date: [TO CONFIRM]", 8.5)
    _replace_cell(sig.cell(1, 0), "Signed (1st Supervisor): [SIGNATURE REQUIRED]", 8.5)
    _replace_cell(sig.cell(1, 1), "Date: [TO CONFIRM]", 8.5)
    path = OUT / "03_Ethical_Consideration_Declaration_Draft.docx"
    doc.save(path)
    return path


def make_logbook() -> Path:
    src = FYP / "Logbook Cover Sheet.docx"
    doc = Document(src)
    configure_a4(doc, margins=(1.5, 1.5, 1.4, 1.4))
    set_doc_defaults(doc, font="Arial", size=9.5, line=1.0)
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].text = "From: 01/10/2026   To: 30/10/2026   |   Draft started 01/10/2026"
        style_run(doc.paragraphs[2].runs[0], 10, True)
    table = doc.tables[0]
    values = [
        "October 2026",
        "1. FYP guidance and deadline review\n2. Project scope and literature plan\n3. Risk/ethics drafts\n4. Preliminary report draft\n5. Reproducibility and annotation plan",
        "[TO CONFIRM]",
        "[TO CONFIRM]",
        "[TO CONFIRM]",
        "[TO CONFIRM — do not invent meetings or emails]",
    ]
    for i, value in enumerate(values):
        _replace_cell(table.rows[2].cells[i], value, 8.5)
    for row in table.rows[3:]:
        for cell in row.cells:
            _replace_cell(cell, "", 8.0)

    doc.add_page_break()
    doc.add_heading("How to maintain this logbook", level=1)
    add_banner(doc, "This starter records only work actually completed on 1 October 2026. Future prompts are plans and must be replaced by genuine contemporaneous entries.", fill="FCE4D6", color=(156, 63, 28))
    add_bullets(doc, [
        "Start each working session with date, page number, start time, objective and intended output; record finish time and next action.",
        "Record commands, versions, parameters, input/output paths, counts, hashes, screenshots, diagrams and links to source material. Paste or reference long outputs rather than copying them without interpretation.",
        "Write hypotheses, observations, mistakes, failed attempts and reasons for decisions. Do not rewrite the narrative later to make the project look linear.",
        "For meetings: write agenda before the meeting; then decisions, questions, action owner and due date. Do not record a meeting that did not occur.",
        "For literature: record search terms, databases, screening decisions and what changed in the project design. Cite enough detail to retrieve the source.",
        "For results: distinguish raw observation from interpretation. Include units, sample size, metric definition and limitations.",
        "Use ink in the physical logbook, sign/date pages, do not remove pages, and correct errors with a single line. The monthly cover sheet is only a summary of the attached pages.",
    ], size=10)

    doc.add_heading("Entry 001 — Guidance review and project initiation", level=1)
    add_table(
        doc,
        ["Field", "Contemporaneous record"],
        [
            ["Date / time", "01/10/2026. Start and finish times: [student to enter from actual work session]."],
            ["Objective", "Understand the 2026/27 FYP submission rules and convert the SMA knowledge-graph blueprint into a truthful future work plan."],
            ["Sources reviewed", "A Guide to FYP Logbook; Student Handbook FYP 2026–27; FYP_timeline.png; Project writing and plagiarism printing 2017; risk, ethics, cover-sheet and preliminary-report templates; project blueprint PDF; repository handoff and reproduction notes."],
            ["Work completed", "Confirmed the preliminary/interim/poster/final deliverables and their weightings. Identified the 28/30 April deadline conflict. Reviewed template fields and page limits. Audited the proposed pipeline against repository evidence. Classified the four red modules as planned innovations rather than completed work. Drafted risk and ethical controls and a forward project schedule."],
            ["Key decision", "No retrospective weekly entries will be created. Existing project materials are treated as a technical blueprint/baseline; from today onward, only tasks actually performed by the student will be recorded as completed."],
            ["Open questions", "Confirm official English name, both student IDs, degree programme, course code, first supervisor, signature process, allowed GenAI declaration, and the authoritative final deadline on Moodle/FYP System."],
            ["Next action", "Create a literature matrix and define the biomedical entity/relation schema; arrange the first real supervisor meeting; update this logbook after each session."],
        ],
        widths=[3.2, 14.6],
        font_size=9.2,
    )

    doc.add_page_break()
    doc.add_heading("Technical planning notes recorded on 01/10/2026", level=2)
    add_para(doc, "Proposed pipeline: PubMed/Open Targets acquisition → BERTopic + PubMedBERT topic analysis → LLM entity/relation extraction → biomedical schema validation → dictionary mapping and embedding-based entity alignment → triple aggregation and conflict detection → Neo4j import/visualisation → GraphRAG and evidence evaluation.", size=10)
    add_para(doc, "Innovation classification from the red markings in the project blueprint:", size=10)
    add_numbered(doc, [
        "Evidence validation: align generated evidence spans back to the PubMed title/abstract and route weak matches to review.",
        "GraphRAG: retrieve Neo4j neighbourhoods and PMID-backed evidence to build a structured, grounded answer context.",
        "Citation validation: check that every generated citation is a real identifier that was actually retrieved.",
        "Claim–evidence validation: decompose an answer into claims and judge whether the cited evidence supports each claim.",
    ], size=9.8)
    add_para(doc, "Evaluation decision: create a 400-item human-reviewed set (300 random + 100 challenge), with 80 independently double-reviewed items. Use the random set for the main precision/support estimate; report challenge analysis separately. Do not claim recall from a prediction-sampled set.", size=10)

    doc.add_heading("Planned prompts for future entries — not evidence of completed work", level=1)
    add_banner(doc, "Delete each prompt only after replacing it with what actually happened. If a task is postponed or fails, record that honestly.", fill="FFF2CC", color=(127, 96, 0))
    prompts = [
        ["02–04 Oct", "Confirm identifiers, supervisor/signature process and official deadline; document actual contact and response."],
        ["05–11 Oct", "Record PubMed/Open Targets/SMA/KG/biomedical NLP search strings, papers read, critique and resulting schema decisions."],
        ["12–18 Oct", "Record environment setup, acquisition commands, data counts, hashes, errors, fixes and validation results."],
        ["19–25 Oct", "Record BERTopic/PubMedBERT setup, parameter choices, cluster interpretation, LLM pilot design and error cases."],
        ["26–30 Oct", "Record report revisions, supervisor feedback, signatures, page/format checks, plagiarism check and upload receipts."],
    ]
    add_table(doc, ["Window", "What the actual entry should capture"], prompts, widths=[3.2, 14.6], font_size=9.4)
    doc.add_heading("Reusable daily entry template", level=2)
    add_table(
        doc,
        ["Section", "Write during or immediately after the work"],
        [
            ["Date / page / time", "Date: ____  Page: ____  Start: ____  Finish: ____  Total: ____"],
            ["Objective", "What specific question or output will be completed?"],
            ["Inputs / setup", "Files, data version/hash, environment, software/model version, command/parameters."],
            ["Actions", "Chronological steps. Include failures and changes rather than only the final successful command."],
            ["Observations", "Counts, tables, plots, examples, unexpected behaviour, uncertainty."],
            ["Interpretation / decision", "What does the evidence mean? What options were considered and why was one selected?"],
            ["Next action", "Concrete task, owner, deadline and prerequisite."],
            ["Signature", "Student initials/signature: ____  Date: ____"],
        ],
        widths=[3.2, 14.6],
        font_size=9.2,
    )
    # Word requires a paragraph after a final table; make it tiny so it does not
    # create a visually blank extra page in the supplied branded template.
    tail = doc.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    style_run(tail.add_run(" "), 1)
    path = OUT / "04_First_Logbook_Starter_October_2026.docx"
    doc.save(path)
    return path


def add_citation_para(doc, text: str, size=8.4):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(text), size)


def make_preliminary() -> Path:
    src = TEMP / "preliminary_template_converted.docx"
    doc = Document(src)
    configure_a4(doc, margins=(1.55, 1.45, 1.55, 1.55))
    set_doc_defaults(doc, font="Arial", size=9.4, line=1.02)
    style_preliminary_headings(doc)

    table = doc.tables[0]
    _replace_cell(table.cell(0, 1), STUDENT, 8.5)
    _replace_cell(table.cell(0, 3), "Not applicable", 8.5)
    _replace_cell(table.cell(1, 1), "[UoG ID TO CONFIRM]", 8.5)
    _replace_cell(table.cell(1, 3), PROJECT_TITLE, 8.0)
    _replace_cell(table.cell(2, 1), "[UESTC ID TO CONFIRM]", 8.5)
    _replace_cell(table.cell(2, 3), SUPERVISOR, 8.5)
    _replace_cell(table.cell(3, 1), "[DEGREE PROGRAMME TO CONFIRM]", 8.2)
    _replace_cell(table.cell(4, 1), "2026–27", 8.5)
    # Preserve the originality declaration and add an explicit signature placeholder.
    table.cell(3, 3).paragraphs[-1].add_run("\nSigned (Student): [SIGNATURE REQUIRED]")

    doc.paragraphs[0].text = "GC-UESTCHN4008P Project Specifications and Preliminary Report"
    style_run(doc.paragraphs[0].runs[0], 11, True, (31, 78, 121))
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[1].text = "Course code: [TO CONFIRM]    Academic year: 2026–27"
    style_run(doc.paragraphs[1].runs[0], 9, True)
    doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Remove the template instruction block while retaining the university-branded first page.
    for p in list(doc.paragraphs[5:]):
        remove_paragraph(p)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    style_run(p.add_run(PROJECT_TITLE), 18, True, (31, 78, 121))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("Project Specifications and Preliminary Report"), 12, True, (47, 84, 117))
    add_banner(doc, "Submission target: Friday 30 October 2026, 23:59 CST — confirm against the current Moodle/FYP System brief before upload.", fill="E2F0D9", color=(55, 86, 35))

    doc.add_page_break()
    doc.add_heading("1. Project Description", level=1)
    add_para(doc, "Spinal muscular atrophy (SMA) is an inherited neuromuscular disorder in which reduced survival motor neuron protein leads to progressive motor-neuron loss. The biomedical evidence surrounding SMA spans molecular mechanisms, genotypes, phenotypes, therapies and outcomes, but much of it remains distributed across literature and heterogeneous databases. This project will design and evaluate a reproducible software pipeline that converts public PubMed abstracts and Open Targets associations into an evidence-grounded SMA knowledge graph.", size=9.4)
    add_para(doc, "The proposed pipeline will acquire and version source data, identify literature topics, extract typed biomedical triples with a large language model (LLM), validate them against a restricted schema, normalise and align synonymous entities, aggregate repeated evidence, detect relation conflicts, and import the result into Neo4j for analysis and visualisation. The graph will then support evidence-grounded retrieval. The intended research contribution is not a clinical decision system: it is a transparent engineering study of how provenance, human evaluation and graph-based retrieval can reduce unsupported biomedical statements.", size=9.4)

    doc.add_heading("2. Measurable Outcomes", level=1)
    add_table(
        doc,
        ["Task / target", "Measurable acceptance criterion", "Tangible output"],
        [
            ["Data acquisition", "Versioned PubMed and Open Targets inputs; no malformed JSON; counts and SHA-256 recorded", "Raw JSONL, manifest and validation report"],
            ["Topic analysis", "Reproducible document embeddings/clusters with interpretable topic terms and parameter record", "Clustered abstracts and topic summary"],
            ["Triple extraction", "Typed subject–relation–object records with PMID, evidence span and validation status", "Canonical JSONL and rejected-record log"],
            ["Fusion and graph", "Normalised entities, aggregated evidence, conflict flags; successful Neo4j import with topology checks", "Neo4j graph, metrics and HTML viewer"],
            ["Human evaluation", "400 reviewed predictions: 300 random main set, 100 challenge set; 80 independent second reviews", "Gold workbook, precision/support metrics, agreement and error analysis"],
            ["Innovation modules", "Evidence alignment, GraphRAG, citation validation and claim–evidence validation tested against explicit baselines", "Software, tests, QA set and ablation tables"],
        ],
        widths=[3.7, 8.1, 5.5],
        font_size=8.2,
    )
    add_para(doc, "The project is primarily a software and theoretical/experimental research output. No bespoke hardware is required. Success will be judged by reproducibility, extraction factuality, traceable provenance, graph integrity and the measured grounding of generated answers—not by graph size alone.", size=9.2)

    doc.add_page_break()
    doc.add_heading("3. Technical Background", level=1)
    doc.add_heading("3.1 Topic Basis and Significance", level=2)
    add_para(doc, "SMA was linked to the survival motor neuron (SMN) gene region by Lefebvre et al. [1]. Subsequent therapeutic development has produced mechanistically distinct interventions: nusinersen modifies SMN2 splicing [2], onasemnogene abeparvovec supplies a functional SMN1 transgene [3], and risdiplam is an oral SMN2 splicing modifier [4]. The field therefore contains relations among genes, variants, molecular functions, drugs, phenotypes, patient subgroups and measured outcomes. These relations are precisely the kind of heterogeneous evidence that benefits from an explicit graph representation.", size=9.5)
    add_para(doc, "The research problem is not lack of publications but the difficulty of integrating and auditing them. A researcher may need to ask which genes, pathways or therapies are connected to a phenotype, which papers support the connection, whether different studies contradict one another, and whether a generated summary cites evidence that actually entails its claims. A conventional keyword search returns documents; an evidence-grounded knowledge graph aims to expose typed relationships while retaining source provenance.", size=9.5)
    add_para(doc, "Open Targets provides public, systematically integrated target–disease evidence and programmatic access, making it a suitable structured complement to PubMed text [5]. PubMed supplies breadth and explicit bibliographic identifiers, whereas Open Targets provides curated/scored association records. Combining them allows the project to compare literature-derived relations with an external structured source without using private patient data.", size=9.5)
    doc.add_heading("3.2 Why a knowledge graph", level=2)
    add_para(doc, "A knowledge graph represents entities as nodes and typed relations as edges. In biomedicine, this model supports integration across terminology and data sources, neighbourhood exploration, provenance-aware aggregation and downstream prediction or retrieval [6]. For SMA, a graph can connect Disease, Gene, Protein, Drug, Phenotype, Pathway and Variant entities while attaching PMID evidence and confidence to each relation. Neo4j is proposed because its property-graph model can store these attributes and Cypher can express multi-hop questions clearly.", size=9.5)
    add_para(doc, "Graph construction nevertheless creates risks. Entity names are ambiguous; the same concept may appear under multiple forms; relation polarity may conflict; and high-confidence model output can still be unsupported. The design must therefore treat schema validation, normalisation, provenance and human review as first-class components rather than post-processing conveniences.", size=9.5)

    doc.add_page_break()
    doc.add_heading("3.3 Literature Review: From text to structured evidence", level=2)
    add_para(doc, "Biomedical relation extraction converts unstructured sentences into semantic links between entities. Earlier systems relied on patterns and syntactic rules; neural and transformer-based approaches improved contextual representation but still face overlapping relations, long-range context, domain terminology and limited labelled data. Literature-scale systems such as SemRep demonstrate the value of interpretable semantic relations, while also illustrating a precision–recall trade-off [7]. More recent graph and attention models address cross-sentence dependencies, but measured performance remains task- and corpus-dependent.", size=9.4)
    add_para(doc, "Domain-specific pre-training is relevant because biomedical vocabulary and sense distributions differ from general web text. PubMedBERT was trained from scratch on biomedical text and achieved strong results across biomedical NLP tasks [8]. In this project it will be used as an embedding model for topic analysis and type-constrained entity alignment, not as proof that two names are identical. Candidate pairs must share an entity type and pass a documented similarity threshold; dictionary mappings take priority where explicit aliases are available.", size=9.4)
    add_para(doc, "BERTopic embeds documents, clusters them and derives interpretable topic descriptors using class-based TF–IDF [9]. It can reveal dominant themes in a large SMA corpus and support topic-aware inspection, but its clusters are sensitive to the embedding model, dimensionality reduction and density parameters. The project will therefore preserve parameters and seeds, inspect representative abstracts and treat topics as exploratory structure rather than ground truth.", size=9.4)
    add_para(doc, "LLMs can extract flexible typed relations with less task-specific training data, but they may generate nonexistent entities, reverse relation direction, overgeneralise evidence or output invalid structures. A restricted biomedical schema, JSON validation, confidence decomposition, exact source retention and rejected-record logs are necessary controls. Crucially, automated confidence scores cannot replace a human-reviewed evaluation set.", size=9.4)
    doc.add_heading("3.4 Evaluation requirement", level=2)
    add_para(doc, "The proposed 400-item annotation set will contain 300 reproducibly sampled predictions for the main factuality/precision estimate and 100 deliberately difficult cases for qualitative error analysis. Eighty items will receive independent second review to measure agreement. Labels will separately assess entity boundaries/types, relation, direction and evidence support. Because the sample begins from predicted triples, it can estimate precision but not recall; a recall study would require exhaustive annotation of all relations in a separate sample of complete abstracts.", size=9.4)

    doc.add_page_break()
    doc.add_heading("3.5 Proposed System Architecture", level=2)
    add_numbered(doc, [
        "Acquisition and versioning. Query PubMed for SMA literature and retrieve Open Targets SMA associations. Store immutable raw files, retrieval dates, query strings, counts, hashes and validation summaries.",
        "Topic modelling. Generate PubMedBERT-based document embeddings and BERTopic clusters. Review topic words and representative abstracts to identify coverage gaps and guide analysis.",
        "LLM extraction. Prompt the model to return typed triples with PMID and evidence text. Process in checkpointed chunks with retry/backoff; preserve raw responses and model/version parameters.",
        "Biomedical validation. Normalise relation names against a controlled schema; validate entity types and required fields; reject invalid records rather than silently repairing meaning.",
        "Entity and relation fusion. Apply curated alias dictionaries, then type-constrained semantic alignment. Aggregate duplicate evidence, combine confidence transparently and mark positive/negative relation conflicts for review.",
        "Graph construction and analysis. Import entities and relations into Neo4j, preserving provenance. Verify node/edge counts, isolated nodes, duplicate relationships and query behaviour; build a local viewer and analytics tables.",
        "Evaluation and evidence-grounded retrieval. Complete human review, quantify extraction quality, then evaluate graph retrieval and generated answers against citation and claim-support checks.",
    ], size=9.2)
    doc.add_heading("3.6 Engineering controls and reproducibility", level=2)
    add_para(doc, "Each pipeline stage will write to a dated run directory before promotion to a canonical output. A manifest will record inputs, outputs, counts, hashes, configuration and validation status. API keys and database credentials will remain in ignored local environment files. Automated tests will cover schema normalisation, unsafe graph labels, deterministic aggregation and validation gates. A stage will be reported as complete only when source code, tests, a reproducible run artefact and measured output all exist.", size=9.4)
    add_para(doc, "Potential failure modes include API outages, rate limits, malformed model output, model drift, embedding download failure, destructive graph re-imports and hidden data leakage. Mitigations are small pilot runs, retries with bounded budgets, checkpointing, pinned model names, local caches, backups, explicit database scope and validation before promotion. These controls make the methodology defensible even if individual model calls are nondeterministic.", size=9.4)
    doc.add_heading("3.7 Data protection, ethics and scope", level=2)
    add_para(doc, "Only public titles, abstracts and public association records are planned. No participant recruitment, patient record, biological sample or intervention is involved. The graph is a research artefact and must not be presented as clinical advice. Source licences and attribution will be respected, full copyrighted papers will not be redistributed, and any permitted generative-AI assistance in writing or code will be disclosed and verified by the student.", size=9.4)

    doc.add_page_break()
    doc.add_heading("3.8 Research Status, Development Trend and Gap", level=2)
    add_para(doc, "The development trend is moving from isolated extraction models toward systems that integrate heterogeneous sources, retain provenance and connect retrieval to generation. Retrieval-augmented generation (RAG) combines a language model with explicit external memory and can improve factuality and provenance compared with a purely parametric model [10]. GraphRAG extends retrieval by organising entities and evidence as a graph, enabling local neighbourhood questions and global corpus-level synthesis [11].", size=9.4)
    add_para(doc, "However, retrieval alone does not guarantee a correct answer. A system can retrieve the right paper but cite the wrong identifier, or cite a real paper whose abstract does not support the generated claim. Biomedical use raises the standard further because plausible but unsupported statements can mislead. The gap addressed here is therefore end-to-end traceability: an extracted relation should be linked to an exact source span; a retrieved context should expose PMIDs; a generated citation should correspond to an actually retrieved record; and each answer claim should be checked against evidence.", size=9.4)
    doc.add_heading("3.9 Planned Original Contribution / Innovation", level=2)
    innovations = [
        ("Innovation 1 — Evidence validation", "Align every extracted evidence span back to its PubMed title/abstract using exact and fuzzy/semantic checks. Strong matches are accepted; weak or cross-sentence matches are routed for review. Evaluate coverage, false acceptance and false rejection on the human set."),
        ("Innovation 2 — GraphRAG", "Retrieve a bounded Neo4j neighbourhood plus PMID-backed evidence, construct structured context and generate answers with an explicit refusal when evidence is insufficient. Compare with text-only retrieval on a fixed QA set."),
        ("Innovation 3 — Citation validation", "Verify that citations are syntactically valid, exist in the source corpus and were present in the retrieved context. Measure missing, invalid and mismatched citation rates."),
        ("Innovation 4 — Claim–evidence validation", "Decompose each answer into atomic claims and classify whether retrieved evidence supports, partially supports, contradicts or cannot resolve each claim. Report supported-claim rate and typical failure modes."),
    ]
    for title, body in innovations:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        style_run(p.add_run(title + ". "), 9.4, True, (192, 0, 0))
        style_run(p.add_run(body), 9.4)
    add_banner(doc, "Status rule: these four modules are planned innovations. They must not be described as implemented contributions until code, focused tests, reproducible run artefacts and measured evaluation results exist.", fill="FCE4D6", color=(156, 63, 28))
    doc.add_heading("3.10 Research Questions", level=2)
    add_bullets(doc, [
        "RQ1: How accurately can a schema-constrained LLM pipeline extract PMID-supported SMA relations from public abstracts?",
        "RQ2: How do dictionary mapping and type-constrained semantic alignment affect duplicate reduction and erroneous merges?",
        "RQ3: Can explicit evidence, citation and claim validation improve the grounding of graph-based answers over an unvalidated retrieval baseline?",
        "RQ4: Which errors remain after validation, and what trade-offs arise between coverage, precision, cost and review workload?",
    ], size=9.2)

    doc.add_page_break()
    doc.add_heading("4. References", level=1)
    refs = [
        "[1] S. Lefebvre et al., “Identification and characterization of a spinal muscular atrophy-determining gene,” Cell, vol. 80, no. 1, pp. 155–165, 1995. doi:10.1016/0092-8674(95)90460-3.",
        "[2] R. S. Finkel et al., “Nusinersen versus sham control in infantile-onset spinal muscular atrophy,” N. Engl. J. Med., vol. 377, no. 18, pp. 1723–1732, 2017. doi:10.1056/NEJMoa1702752.",
        "[3] J. W. Day et al., “Onasemnogene abeparvovec gene therapy for symptomatic infantile-onset spinal muscular atrophy … (STR1VE),” Lancet Neurol., vol. 20, no. 4, pp. 284–293, 2021. doi:10.1016/S1474-4422(21)00001-6.",
        "[4] E. Mercuri et al., “Safety and efficacy of once-daily risdiplam in type 2 and non-ambulant type 3 spinal muscular atrophy (SUNFISH part 2),” Lancet Neurol., vol. 21, no. 1, pp. 42–52, 2022. doi:10.1016/S1474-4422(21)00367-7.",
        "[5] D. Ochoa et al., “Open Targets Platform: supporting systematic drug–target identification and prioritisation,” Nucleic Acids Res., vol. 49, no. D1, pp. D1302–D1310, 2021. doi:10.1093/nar/gkaa1027.",
        "[6] C. Nicholson and C. S. Greene, “Constructing knowledge graphs and their biomedical applications,” Comput. Struct. Biotechnol. J., vol. 18, pp. 1414–1428, 2020. doi:10.1016/j.csbj.2020.05.017.",
        "[7] H. Kilicoglu et al., “Broad-coverage biomedical relation extraction with SemRep,” BMC Bioinformatics, vol. 21, 2020. PMID:32410573.",
        "[8] Y. Gu et al., “Domain-specific language model pretraining for biomedical natural language processing,” ACM Trans. Comput. Healthcare, vol. 3, no. 1, pp. 1–23, 2021. doi:10.1145/3458754.",
        "[9] M. Grootendorst, “BERTopic: Neural topic modeling with a class-based TF-IDF procedure,” arXiv:2203.05794, 2022. doi:10.48550/arXiv.2203.05794.",
        "[10] P. Lewis et al., “Retrieval-augmented generation for knowledge-intensive NLP tasks,” Adv. Neural Inf. Process. Syst., vol. 33, 2020.",
        "[11] D. Edge et al., “From local to global: A graph RAG approach to query-focused summarization,” arXiv:2404.16130, 2024. doi:10.48550/arXiv.2404.16130.",
    ]
    for ref in refs:
        add_citation_para(doc, ref, size=8.2)
    doc.add_heading("Reference selection note", level=2)
    add_para(doc, "The final literature review will expand this preliminary list through a documented database search and will verify all bibliographic metadata against PubMed, Crossref or the publisher record. Web pages will be used only where a stable primary technical source is unavailable.", size=8.8)

    doc.add_page_break()
    doc.add_heading("5. Work Plan", level=1)
    add_para(doc, "The plan begins on 1 October 2026. Dates below are future targets, not reconstructed historical activity. Each milestone will be marked complete only after a validated output and a contemporaneous logbook entry exist.", size=9.2)
    add_table(
        doc,
        ["Period", "Planned work", "Milestone / evidence"],
        [
            ["01–30 Oct 2026", "Scope, literature, schema, risk/ethics; acquisition and small extraction pilot; preliminary submission", "Signed forms; report; first logbook; versioned inputs and pilot error analysis"],
            ["Nov 2026", "Full LLM extraction; schema validation; dictionary mapping and semantic alignment", "Canonical/rejected triples; tests; manifests; threshold study"],
            ["Dec 2026", "Aggregation, conflict detection, Neo4j import, analytics, visualisation; interim draft", "Fused graph; conflict set; topology metrics; viewer"],
            ["01–08 Jan 2027", "Reproduce initial results and submit interim report and second logbook", "Interim package and revised Gantt"],
            ["11–22 Jan 2027", "Interim presentation and feedback response", "Presentation, Q&A record and action list"],
            ["23 Jan–14 Feb", "Human review of 400 triples; 80 double reviews", "Gold labels, agreement, precision/support and errors"],
            ["15 Feb–07 Mar", "Innovation 1: evidence-span validation", "Implementation, tests and validation metrics"],
            ["08–21 Mar", "Innovation 2: GraphRAG", "Retriever/generator, QA set and baseline comparison"],
            ["22 Mar–04 Apr", "Innovations 3–4: citation and claim–evidence validation", "Grounding metrics and ablation results"],
            ["05–16 Apr", "Final reruns, figures, poster and oral rehearsal", "Frozen results and poster submission"],
            ["19–21 Apr", "Poster-based oral assessment", "Panel feedback and report actions"],
            ["22–28 Apr", "Final dissertation, third logbook and code archive; internal deadline", "Complete submission package by 28 April"],
            ["29–30 Apr", "Contingency only if official systems confirm 30 April", "Upload receipt and archive"],
        ],
        widths=[3.1, 8.7, 5.7],
        font_size=8.0,
    )
    doc.add_heading("Project Outline", level=2)
    add_para(doc, "The critical path is acquisition → extraction → validation/fusion → graph → human evaluation → evidence-grounded retrieval and answer validation. Human annotation is scheduled before the innovation evaluation so that thresholds and conclusions are anchored in independently reviewed evidence. The four innovation modules will be reported together as Original Contribution, but each will retain a separate completion status and metric. The internal final deadline is 28 April 2027 because the supplied timeline and handbook disagree; the live Moodle/FYP System notice will determine the formal deadline.", size=9.2)
    add_heading = doc.add_paragraph()
    style_run(add_heading.add_run("Resources: "), 9.2, True)
    style_run(add_heading.add_run("Existing computer, Python/Conda environment, public data, Neo4j and approved API access. No component purchase or reimbursement request is currently required."), 9.2)
    p = doc.add_paragraph()
    style_run(p.add_run("Risk Assessment: "), 9.2, True)
    style_run(p.add_run("A completed draft is supplied separately and must be checked and signed by the student and first supervisor before submission."), 9.2)

    path = OUT / "05_GCU_Specification_and_Preliminary_Report_Draft.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    outputs = [make_roadmap(), make_risk(), make_ethics(), make_logbook(), make_preliminary()]
    for p in outputs:
        print(f"{p.name}\t{p.stat().st_size}")


if __name__ == "__main__":
    main()

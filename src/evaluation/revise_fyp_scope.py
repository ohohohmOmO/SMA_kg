"""Align five supplied FYP drafts with the measured core project; preserve originals."""
import argparse
import hashlib
import json
import shutil
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt


def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def paragraph_text(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]: run.text = ""
    else: paragraph.add_run(text)


def cell_text(cell, text):
    paragraph_text(cell.paragraphs[0], text)
    for p in cell.paragraphs[1:]: p._element.getparent().remove(p._element)


def after(paragraph, text, style="Normal"):
    element = OxmlElement("w:p"); paragraph._p.addnext(element)
    result = Paragraph(element, paragraph._parent); result.style = style
    result.add_run(text)
    return result


def preliminary(doc):
    replacements = {
        10: "The pipeline versions source data, analyses literature topics, extracts typed biomedical triples, validates their schema, applies curated dictionaries and conservative typed identity normalisation, aggregates evidence and flags conflicts. It imports a versioned Neo4j graph and supports source-backed visual inspection. The contribution is an auditable engineering study of extraction support, identity errors and evidence screening; graph-based answer generation is future work.",
        12: "The project is a software and experimental research output. Success is assessed through reproducibility, source support, provenance preservation, typed graph integrity and controlled screening results. Strict support, partial support and evidence-span adequacy are separate measures. Graph size and heuristic confidence do not measure factual accuracy.",
        17: "The integration problem is to identify which typed entities and relations occur in the stored literature, expose their supporting PMIDs and source sentences, and distinguish conflicting or conditional statements. A keyword search returns documents; this graph exposes directed relationships while retaining original predictions and their provenance. Generated-answer validation is outside the core evaluation.",
        20: "A knowledge graph represents entities as nodes and typed relations as edges, supporting neighbourhood exploration and provenance-aware aggregation [6]. The implemented schema contains Disease, Gene, Protein, Drug, Phenotype and Variant. Pathway is not currently an accepted entity type. Neo4j stores typed identities, evidence records and review status; public structured associations retain their source identifiers separately.",
        25: "Biomedical pre-training provides useful domain representations [8], but sentence similarity cannot establish entity identity. The historical PubMedBERT embedding alignment merged distinct genes and clinical subtypes. The repaired design uses type plus name throughout and automatically normalises only equivalent case/spacing forms after guarded dictionary mapping. Embedding similarity is reserved for review proposals. Different official gene identifiers and subtype qualifiers must remain distinct.",
        29: "The supplied 400-item set has been completed: 300 random predictions for the main support estimate and 100 challenge cases reported separately. The student confirmed that all labels were human-assigned and ChatGPT only assembled the workbook. Integrated support and evidence adequacy are available; component type/direction fields and independent second reviews are empty. No agreement statistic is claimed. Prediction sampling cannot establish extraction recall or F1. Further independent review is optional additional evidence, not a completed commitment.",
        34: "3. LLM extraction. Retain model parameters, checkpointed predictions, PMID and original evidence. The revised prompt requires complete exact source sentences, distinguishes genes/proteins/variants and prohibits dropping material conditions. Prompt revisions are not credited with an accuracy gain until newly extracted predictions receive appropriate evaluation.",
        35: "4. Evidence and assertion screening. Check source offsets, complete context, type contradictions, directed relation and retained conditions. Exact quotes do not prove entailment. A separate model-assisted screen receives the original assertion and source text without human labels; invalid, ambiguous and incomplete outputs fail closed and remain reviewable.",
        36: "5. Entity and relation fusion. Apply guarded curated aliases and typed orthographic identity. Preserve digits, variants and subtype qualifiers; semantic similarity does not automatically merge entities. Aggregate repeated evidence and mark potential polarity conflicts. Compare raw, dictionary, repaired identity and the preserved historical semantic baseline.",
        37: "6. Graph construction and analysis. Import a versioned typed graph without clearing the historical graph. Keep literature and Open Targets identities and relation sources separate. Reconcile every node, relationship and evidence property with the validated inputs; verify repeat-import idempotence, topology and the offline source explorer.",
        38: "7. Evaluation and discussion. Report human strict/partial support, errors, controlled fusion results and paired evidence/context screening. Present selective support precision beside retention, false acceptance and strictly supported candidates routed to review. Use the viewer for inspection; do not claim a completed GraphRAG QA system.",
        47: "The core gap addressed here is auditable construction rather than generated-answer grounding: original predictions must remain linked to saved source spans; identity normalisation must preserve medically distinct entities; and screening must reveal conditions without silently treating candidates as verified facts. RAG and GraphRAG [10,11] motivate future extensions, but generated-answer citation checking and atomic claim validation are not current deliverables.",
        48: "3.9 Original Contributions and Current Evidence",
        49: "Contribution 1 - Typed identity repair and controlled fusion. Correct the name-only mapping overwrite, preserve SMN1/SMN2 and SMA subtypes, retain all 18,288 evidence records, and compare structural effects against the historical unsafe baseline. The repaired graph contains 13,001 literature edges. Compression is not mapping accuracy; separate mapping review remains available.",
        50: "Contribution 2 - Evidence traceability and assertion screening. Locate spans with source offsets; retain complete contexts and review reasons; compare the previous evidence gate with type/context rules, label-blind model screening and their combination. Fuzzy suggestions are review-only. Original candidates and human labels are never rewritten or replaced.",
        51: "Measured limitation - On the 202-prediction random internal test, combined screening retains 30 candidates with 46.7% strict support, versus 40.7% for the previous gate. It retains only 14 of 55 strict positives. The rule-only screen performs poorly. These retrospective results support review prioritisation, not sufficient extraction quality or a full-graph accuracy claim.",
        52: "Future work - GraphRAG answer generation, generated-answer citation validation and atomic claim-to-evidence validation remain unimplemented extensions. The limited original-assertion model screen above does not complete those answer-level modules. Their scope requires a separate QA dataset and evaluation design.",
        54: "RQ1: What strict and partial source support do the original schema-constrained SMA predictions achieve, and which condition, type and evidence errors recur?",
        55: "RQ2: How do raw naming, dictionary normalisation and repaired typed identity affect graph structure and provenance, compared with historical semantic fusion?",
        56: "RQ3: Can evidence/context/type screening enrich strict support among retained original predictions, and what retention, false-acceptance and strict-positive-loss trade-offs arise?",
        57: "Evaluation boundary: keep random and challenge samples separate; distinguish source location from entailment; report the internal retrospective design and absent independent agreement or exhaustive recall.",
        75: "The schedule starts from the verified engineering state on 2 October 2026. Existing artifacts are evidence of completed automated work, not a reconstructed account of personal study or supervisor meetings. Future tasks concentrate on interpretation, reproducibility, literature critique and assessed writing; signatures and authentic logbook entries remain the student's responsibility.",
        77: "The critical path is fixed scope and metrics -> completed human-label integration -> controlled fusion -> evaluated evidence/assertion screening -> results and discussion. The four requested engineering/reporting repairs now have dated evidence. Remaining work deepens analysis and presentation; GraphRAG and answer citation/claim modules are future work. The internal final deadline remains 28 April 2027, subject to the current Moodle/FYP System notice.",
    }
    for i, text in replacements.items(): paragraph_text(doc.paragraphs[i], text)
    cells = {
        (2,4,1): "Typed identity, source-specific relations and complete provenance; current online graph reconciled against inputs",
        (2,4,2): "Versioned Neo4j graph, acceptance JSON, typed metrics and source explorer",
        (2,5,1): "400 human-labelled original predictions: 300 random + 100 challenge; integrated support/span checks; no second-review agreement supplied",
        (2,5,2): "Preserved workbook, support intervals, error analysis and explicit measurement limits",
        (2,6,0): "Controlled quality improvement",
        (2,6,1): "Evidence, context/type and model-assisted screens compared on unchanged original predictions; precision and retention reported together",
        (2,6,2): "Frozen protocol, label-blind requests, decisions, tests and paired comparison tables",
        (3,0,0): "Status rule: measured contributions require code, regression checks, dated artifacts and evaluation. GraphRAG and answer-level citation/claim validation are future work. No model decision is relabelled as human verification.",
    }
    plan = [
        ("01-30 Oct 2026", "Review repaired pipeline and measured limits; refine literature, scope, ethics/risk and preliminary report", "Repaired run artifacts; current specification; signed forms and genuine logbook"),
        ("Nov 2026", "Reproduce controlled fusion/screening; examine failure cases and source/type ambiguity", "Hash-checked rerun; paired tables; critical case analysis"),
        ("Dec 2026", "Verify demonstration and database queries; integrate methods and interim results", "Typed graph acceptance; viewer checks; interim draft"),
        ("01-08 Jan 2027", "Submit interim report and second logbook; revise schedule from actual progress", "Interim package and revised Gantt"),
        ("11-22 Jan 2027", "Interim presentation and response to actual feedback", "Presentation and genuine feedback/action record"),
        ("23 Jan-14 Feb", "Analyse existing 400 labels; optionally review the fixed 30 mappings", "Support/error analysis; mapping quality only if new judgments supplied"),
        ("15 Feb-07 Mar", "Strengthen discussion of screening failures, coverage loss and uncertainty", "Case analysis and precision/retention figures"),
        ("08-21 Mar", "Complete literature critique, methods and results integration", "Dissertation chapters and reproducibility appendix"),
        ("22 Mar-04 Apr", "Check conclusions against objectives; audit references, limits and source attribution", "Complete draft and traceable claims"),
        ("05-16 Apr", "Freeze results, prepare poster and oral rehearsal", "Final figures and poster submission"),
        ("19-21 Apr", "Poster-based oral assessment", "Actual panel feedback and report actions"),
        ("22-28 Apr", "Final dissertation, third logbook and code archive; internal deadline", "Complete submission package by 28 April"),
        ("29-30 Apr", "Contingency only if official systems confirm 30 April", "Upload receipt and archive"),
    ]
    for j, row in enumerate(plan, 1):
        for c, value in enumerate(row): cells[(4,j,c)] = value
    for (t,r,c), text in cells.items(): cell_text(doc.tables[t].rows[r].cells[c], text)


def roadmap(doc):
    paragraph_text(doc.paragraphs[4], "2. 基于已验证现状的后续执行计划")
    paragraph_text(doc.paragraphs[5], "安排原则：2026-10-02 的实体融合修复、400条人工标签整合、受控筛查和在线数据库验收已有运行材料。后续不重复安排为从零开发任务，重点转向理解方法、分析失败、复现实验和写作。已有产物不能代替学生真实学习记录或导师反馈。")
    paragraph_text(doc.paragraphs[29], "400条标签已由用户确认全部人工标记，ChatGPT仅整理表格。随机主集严格支持26.3%，部分支持47.0%；73.3%为宽松支持率，不能称为完全正确率。学校未硬性要求400条或80条双评；现有独立第二评审为空，不报告一致性。")
    paragraph_text(doc.paragraphs[30], "现有400条无需重填；组件类型/方向字段为空，不推算其准确率。新的融合映射或改写候选需要对应的新判断，不能沿用原预测标签。")
    paragraph_text(doc.paragraphs[31], "Support label：2=直接支持；1=部分支持；0=不支持；U=无法判断。证据片段合格与关系支持分别计算；新模型筛查始终标为模型辅助。")
    table = doc.tables[3]
    rows = {
        1: ("10/01-10/04", "现状与范围", "核对四项修复产物、指标、风险/伦理、提交要求", "运行清单；修订目标；真实日志"),
        2: ("10/05-10/11", "文献与讨论", "审读SMA、实体链接、关系抽取与评价文献；明确误差定义", "文献矩阵；方法选择与局限"),
        3: ("10/12-10/18", "复现与案例", "按冻结输入复现融合与筛查；解释类型和条件错误", "哈希核对；失败案例；保留率分析"),
        4: ("10/19-10/25", "演示与报告", "验证类型化数据库、来源查看和抽取质量讨论", "查询与viewer验收；初稿"),
        6: ("11/01-11/15", "抽取误差研究", "分析400条已有人工标签；不得把宽松支持当完全正确", "严格/部分支持、误差和区间"),
        7: ("11/16-11/30", "融合分析", "比较原始/字典/类型化身份和历史语义基线；新30项审核可选", "结构对照；身份回归；未审指标缺失"),
        12: ("01/23-02/14", "深化现有评价", "解释已有标签与规则/模型失败；必要时补充新判断", "错误分析；无独立双评则不报一致性"),
        13: ("02/15-03/07", "筛查讨论", "比较原证据、条件/类型、模型和联合筛查的精度及保留量", "受控表图；误保留和严格支持送审"),
        14: ("03/08-03/21", "论文整合", "完成文献批判、方法与结果；GraphRAG列为未来工作", "完整章节与复现附录"),
        15: ("03/22-04/04", "结论与引用核对", "检查目标-方法-结果对应；答复生成引用/声明验证不作已实现承诺", "完整报告；文献、范围与局限核验"),
    }
    for r, values in rows.items():
        for c, value in enumerate(values): cell_text(table.rows[r].cells[c], value)
    cell_text(doc.tables[4].rows[3].cells[2], "为何不按embedding相似度自动合并；SMN1/SMN2错并怎样修复；条件筛查为何损失严格支持候选")
    cell_text(doc.tables[4].rows[6].cells[2], "实体身份修复、证据定位、条件/类型/模型筛查、数据库溯源验收；GraphRAG与回答引用/声明验证为未来工作")
    cell_text(doc.tables[6].rows[3].cells[0], "Independent second review")
    cell_text(doc.tables[6].rows[3].cells[1], "未提供")
    cell_text(doc.tables[6].rows[3].cells[2], "可选补充；80条不是硬性承诺")
    cell_text(doc.tables[6].rows[3].cells[3], "无独立标签则不报告Cohen's kappa或一致率")


def risk(doc):
    paragraph_text(doc.paragraphs[10], "Software research on an auditable SMA knowledge graph using public PubMed abstracts and Open Targets records. The work evaluates extraction support, guarded typed identity fusion, evidence/context/type screening, versioned Neo4j import and source-backed visual inspection. GraphRAG answer generation and answer citation/claim validation are future work. No laboratory work, participant recruitment, private clinical records, biological samples or animal work is planned.")
    p = doc.paragraphs[21]
    paragraph_text(p, p.text.replace("Controls: export/backup before imports; use scoped scripts and verify target database and counts before destructive operations.", "Controls: preserve historical graphs; import a new typed version, reconcile every node/edge/evidence property, and activate only after acceptance. Do not automatically clear the database."))


def ethics(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "graph/GraphRAG output" in cell.text:
                    cell_text(cell, "This desk-based software project uses public PubMed titles/abstracts and Open Targets records. It does not recruit participants, access patient records, collect identifiable personal data, use biological samples, conduct interventions or involve animals.\n\nThe main issues are accuracy, provenance, copyright, bias and possible misinterpretation as clinical advice. Controls are to preserve source identifiers and evidence; cite external material; use permitted public data; avoid redistributing copyrighted full articles; protect API credentials; evaluate predictions using the existing 400 human labels; retain uncertainty and conflicts; distinguish partial from strict support; and describe graph/model-assisted screening outputs as research rather than diagnosis or treatment.\n\nIndependent annotation assistance is a project contribution; only a minimal reviewer code is recorded. Any future private clinical data, interviews, patient groups or biological work requires renewed ethical review before collection.")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.size = Pt(10.5)


def logbook(doc):
    # Preserve the actual 1 October planning entry, including its historical scope.
    anchor = doc.paragraphs[24]
    p = after(anchor, "Scope revision on 02/10/2026", "Heading 2")
    p = after(p, "The 01/10 planning notes above are historical and superseded by the core scope: typed identity repair, integrated human evaluation, controlled evidence/assertion screening, versioned database acceptance and source inspection. GraphRAG, generated-answer citation checks and atomic answer-claim validation remain future work. Independent second-review labels have not been supplied.")
    p = after(p, "Assisted engineering session on 02/10/2026: automated work restored 982 SMN2 endpoints, retained all 18,288 evidence records, produced 13,001 literature edges, completed label-blind model screening of 400 unchanged candidates, and reconciled the online typed graph (9,218 nodes; 13,001 literature plus 164 external relations). Old graph counts were preserved. Combined main-test screening retained 30/202 candidates, including 14/55 strict positives; strict support among retained was 46.7%. These results show limited selective enrichment, not adequate full-corpus extraction quality.")
    p = after(p, "Source evidence: artifacts/runs/stage3_identity_repair_2026-10-02, assertion_quality_repair_2026-10-02, fyp_identity_repaired_2026-10-02 and stage4_typed_identity_repair_2026-10-02. Student understanding, personal session times, meetings, supervisor approval and signatures must be recorded from actual events; this entry does not infer them.")
    replacements = {
        1: "Record actual review of the four fixes, scope decision and any genuine supervisor contact; keep unknown identifiers/times unfilled.",
        2: "Record primary-literature searches, critique, error definitions and interpretation of the existing 400-label results.",
        3: "Record controlled reruns, hashes, source/identity checks, failures and the typed database queries actually executed.",
        4: "Record demonstration checks, screening false accepts/rejects, precision-retention discussion and report revisions.",
    }
    for r, text in replacements.items(): cell_text(doc.tables[4].rows[r].cells[1], text)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fyp-root", required=True)
    parser.add_argument("--run-dir", required=True)
    args = parser.parse_args()
    root, out = Path(args.fyp_root), Path(args.run_dir)
    out.mkdir(parents=True, exist_ok=False)
    (out / "originals").mkdir(); (out / "revised").mkdir()
    items = [
        ("00_总览_要求与时间线/01_FYP_Roadmap_and_Submission_Guide.docx", roadmap),
        ("01_首轮提交_2026-10-30/01_交付稿/02_Project_Risk_Assessment_Draft.docx", risk),
        ("01_首轮提交_2026-10-30/01_交付稿/03_Ethical_Consideration_Declaration_Draft.docx", ethics),
        ("01_首轮提交_2026-10-30/01_交付稿/04_First_Logbook_Starter_October_2026.docx", logbook),
        ("01_首轮提交_2026-10-30/01_交付稿/05_GCU_Specification_and_Preliminary_Report_Draft.docx", preliminary),
    ]
    manifest = []
    for relative, revise in items:
        source = root / relative; original = out / "originals" / source.name
        target = out / "revised" / source.name
        shutil.copyfile(source, original)
        doc = Document(original); revise(doc); doc.save(target)
        manifest.append({"source": str(source), "original_snapshot": str(original.resolve()),
            "revised": str(target.resolve()), "before_sha256": digest(original), "after_sha256": digest(target),
            "promoted_to_fyp": False})
    (out / "manifest.json").write_text(json.dumps({"scope": "Construction, conservative identity, evidence/assertion screening and analysis; answer-generation validation is future work",
        "supervisor_approval_claimed": False, "signatures_or_meetings_fabricated": False,
        "historical_log_entry_preserved": True, "documents": manifest}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"documents": len(manifest), "output": str(out.resolve()), "awaiting": "render and visual verification before promotion"}, ensure_ascii=False))


if __name__ == "__main__": main()
